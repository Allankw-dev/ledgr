"""
Bursar report export — CSV and Excel (.xlsx) versions of the per-student
fee ledger for a term. The Excel version uses real formulas (=D{row}-E{row}
for balance, =SUM(...) for the totals row) rather than pre-computed numbers,
so a bursar can open it, tweak a figure, and trust the totals stay correct —
the same reason invoices/payments never store amount_paid as a cached float
computed elsewhere without a way to re-derive it.
"""

import csv
import io
from dataclasses import dataclass
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INK_900 = "16213D"
SURFACE_100 = "F4F2EC"

HEADERS = ["Admission No.", "Student Name", "Class", "Total Billed", "Amount Paid", "Balance", "Status"]


@dataclass
class BursarReportRow:
    admission_number: str
    full_name: str
    class_name: str
    total_amount: Decimal
    amount_paid: Decimal
    status: str


def generate_bursar_report_xlsx(
    school_name: str,
    term_name: str,
    currency: str,
    rows: list[BursarReportRow],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bursar Report"

    money_fmt = f'"{currency}" #,##0.00'

    # Title block
    ws.merge_cells("A1:G1")
    ws["A1"] = f"{school_name} — {term_name} Fee Report"
    ws["A1"].font = Font(bold=True, size=14, color=INK_900)

    header_row = 3
    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=INK_900)
        cell.alignment = Alignment(horizontal="center")

    first_data_row = header_row + 1
    r = first_data_row
    for row in rows:
        ws.cell(row=r, column=1, value=row.admission_number)
        ws.cell(row=r, column=2, value=row.full_name)
        ws.cell(row=r, column=3, value=row.class_name)
        ws.cell(row=r, column=4, value=float(row.total_amount)).number_format = money_fmt
        ws.cell(row=r, column=5, value=float(row.amount_paid)).number_format = money_fmt
        # Real formula, not a pre-computed value — stays correct if a
        # bursar edits a billed/paid figure directly in the sheet.
        ws.cell(row=r, column=6, value=f"=D{r}-E{r}").number_format = money_fmt
        ws.cell(row=r, column=7, value=row.status)
        r += 1
    last_data_row = r - 1

    totals_row = r + 1
    ws.cell(row=totals_row, column=3, value="TOTALS").font = Font(bold=True)
    if last_data_row >= first_data_row:
        for col in (4, 5, 6):
            letter = get_column_letter(col)
            cell = ws.cell(row=totals_row, column=col, value=f"=SUM({letter}{first_data_row}:{letter}{last_data_row})")
            cell.number_format = money_fmt
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=SURFACE_100)

    widths = [16, 26, 14, 16, 16, 16, 16]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_bursar_report_csv(
    school_name: str,
    term_name: str,
    currency: str,
    rows: list[BursarReportRow],
) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    writer.writerow([f"{school_name} — {term_name} Fee Report"])
    writer.writerow([])
    writer.writerow(HEADERS)

    total_billed = Decimal("0")
    total_paid = Decimal("0")
    for row in rows:
        balance = row.total_amount - row.amount_paid
        total_billed += row.total_amount
        total_paid += row.amount_paid
        writer.writerow(
            [
                row.admission_number,
                row.full_name,
                row.class_name,
                f"{row.total_amount:.2f}",
                f"{row.amount_paid:.2f}",
                f"{balance:.2f}",
                row.status,
            ]
        )

    writer.writerow([])
    writer.writerow(
        ["", "", "TOTALS", f"{total_billed:.2f}", f"{total_paid:.2f}", f"{(total_billed - total_paid):.2f}", ""]
    )

    return buffer.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 CSVs cleanly
